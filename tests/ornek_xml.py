"""Uydurma UBL-TR e-Fatura XML'leri üretir ve XML okuyucuyu sınar.

    .venv/Scripts/python tests/ornek_xml.py          # tests/ornekler/ altına yazar ve denetler

Gerçek firma verisi içermez.
"""
import sys
from pathlib import Path

KOK = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(KOK))

NS = ('xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2" '
      'xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2" '
      'xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2"')
XML = '<?xml version="1.0" encoding="UTF-8"?>\n'
BIZ = ("1111111111", "DENEME ALÜMİNYUM A.Ş.", "MERKEZ VD")


def taraf(no, ad, vd, sema="VKN"):
    return (f"<cac:Party><cac:PartyIdentification><cbc:ID schemeID=\"{sema}\">{no}</cbc:ID></cac:PartyIdentification>"
            f"<cac:PartyName><cbc:Name>{ad}</cbc:Name></cac:PartyName><cac:PostalAddress><cbc:StreetName>Deneme Cad."
            f"</cbc:StreetName><cbc:BuildingNumber>1</cbc:BuildingNumber><cbc:CitySubdivisionName>Merkez"
            f"</cbc:CitySubdivisionName><cbc:CityName>Ankara</cbc:CityName><cac:Country><cbc:Name>Türkiye</cbc:Name>"
            f"</cac:Country></cac:PostalAddress><cac:PartyTaxScheme><cac:TaxScheme><cbc:Name>{vd}</cbc:Name>"
            f"</cac:TaxScheme></cac:PartyTaxScheme></cac:Party>")


def vergi(matrah, tutar, oran, kod="0015", ad="KDV", istisna=None, para="TRY"):
    ist = (f"<cbc:TaxExemptionReasonCode>{istisna[0]}</cbc:TaxExemptionReasonCode>"
           f"<cbc:TaxExemptionReason>{istisna[1]}</cbc:TaxExemptionReason>") if istisna else ""
    return (f"<cac:TaxSubtotal><cbc:TaxableAmount currencyID=\"{para}\">{matrah:.2f}</cbc:TaxableAmount>"
            f"<cbc:TaxAmount currencyID=\"{para}\">{tutar:.2f}</cbc:TaxAmount><cbc:Percent>{oran}</cbc:Percent>"
            f"<cac:TaxCategory>{ist}<cac:TaxScheme><cbc:Name>{ad}</cbc:Name><cbc:TaxTypeCode>{kod}</cbc:TaxTypeCode>"
            f"</cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal>")


def kalem(sira, ad, miktar, birim, fiyat, kdv_oran, iskonto=0.0, tevkifat=None, para="TRY", ek_vergi=""):
    brut = round(miktar * fiyat, 2)
    tutar = round(brut - iskonto, 2)
    kdv = round(tutar * kdv_oran / 100, 2)
    isk = (f"<cac:AllowanceCharge><cbc:ChargeIndicator>false</cbc:ChargeIndicator><cbc:MultiplierFactorNumeric>"
           f"{iskonto / brut:.4f}</cbc:MultiplierFactorNumeric><cbc:Amount currencyID=\"{para}\">{iskonto:.2f}</cbc:Amount>"
           f"<cbc:BaseAmount currencyID=\"{para}\">{brut:.2f}</cbc:BaseAmount></cac:AllowanceCharge>") if iskonto else ""
    tev = ""
    if tevkifat:
        kod, pay = tevkifat
        tev = (f"<cac:WithholdingTaxTotal><cbc:TaxAmount currencyID=\"{para}\">{kdv * pay / 10:.2f}</cbc:TaxAmount>"
               + vergi(kdv, kdv * pay / 10, pay * 10, kod, "KDV TEVKIFAT", para=para) + "</cac:WithholdingTaxTotal>")
    return (f"<cac:InvoiceLine><cbc:ID>{sira}</cbc:ID><cbc:InvoicedQuantity unitCode=\"{birim}\">{miktar}"
            f"</cbc:InvoicedQuantity><cbc:LineExtensionAmount currencyID=\"{para}\">{tutar:.2f}</cbc:LineExtensionAmount>"
            f"{isk}<cac:TaxTotal><cbc:TaxAmount currencyID=\"{para}\">{kdv:.2f}</cbc:TaxAmount>"
            f"{vergi(tutar, kdv, kdv_oran, para=para)}{ek_vergi}</cac:TaxTotal>{tev}<cac:Item><cbc:Name>{ad}</cbc:Name>"
            f"<cac:SellersItemIdentification><cbc:ID>STK-{sira}</cbc:ID></cac:SellersItemIdentification></cac:Item>"
            f"<cac:Price><cbc:PriceAmount currencyID=\"{para}\">{fiyat:.2f}</cbc:PriceAmount></cac:Price></cac:InvoiceLine>")


def fatura(no, uuid, tarih, profil, tip, satici, alici, kalemler, vergiler, toplamlar, tevkifat="", para="TRY",
           kur=None, ek=""):
    lext, isk, haric, dahil, odenecek = toplamlar
    kur_x = (f"<cac:PricingExchangeRate><cbc:SourceCurrencyCode>{para}</cbc:SourceCurrencyCode><cbc:TargetCurrencyCode>"
             f"TRY</cbc:TargetCurrencyCode><cbc:CalculationRate>{kur}</cbc:CalculationRate></cac:PricingExchangeRate>"
             ) if kur else ""
    return (f"<Invoice {NS}><cbc:UBLVersionID>2.1</cbc:UBLVersionID><cbc:CustomizationID>TR1.2</cbc:CustomizationID>"
            f"<cbc:ProfileID>{profil}</cbc:ProfileID><cbc:ID>{no}</cbc:ID><cbc:CopyIndicator>false</cbc:CopyIndicator>"
            f"<cbc:UUID>{uuid}</cbc:UUID><cbc:IssueDate>{tarih}</cbc:IssueDate><cbc:IssueTime>10:00:00</cbc:IssueTime>"
            f"<cbc:InvoiceTypeCode>{tip}</cbc:InvoiceTypeCode><cbc:Note>Deneme notu</cbc:Note>"
            f"<cbc:DocumentCurrencyCode>{para}</cbc:DocumentCurrencyCode><cbc:LineCountNumeric>{len(kalemler)}"
            f"</cbc:LineCountNumeric>{ek}"
            f"<cac:AccountingSupplierParty>{taraf(*satici)}</cac:AccountingSupplierParty>"
            f"<cac:AccountingCustomerParty>{taraf(*alici)}</cac:AccountingCustomerParty>"
            f"<cac:PaymentMeans><cbc:PaymentMeansCode>42</cbc:PaymentMeansCode><cbc:PaymentDueDate>2026-10-30"
            f"</cbc:PaymentDueDate><cac:PayeeFinancialAccount><cbc:ID>TR000000000000000000000001</cbc:ID>"
            f"</cac:PayeeFinancialAccount></cac:PaymentMeans>{kur_x}"
            f"<cac:TaxTotal><cbc:TaxAmount currencyID=\"{para}\">0</cbc:TaxAmount>{''.join(vergiler)}</cac:TaxTotal>"
            f"{tevkifat}<cac:LegalMonetaryTotal>"
            f"<cbc:LineExtensionAmount currencyID=\"{para}\">{lext:.2f}</cbc:LineExtensionAmount>"
            f"<cbc:TaxExclusiveAmount currencyID=\"{para}\">{haric:.2f}</cbc:TaxExclusiveAmount>"
            f"<cbc:TaxInclusiveAmount currencyID=\"{para}\">{dahil:.2f}</cbc:TaxInclusiveAmount>"
            f"<cbc:AllowanceTotalAmount currencyID=\"{para}\">{isk:.2f}</cbc:AllowanceTotalAmount>"
            f"<cbc:PayableAmount currencyID=\"{para}\">{odenecek:.2f}</cbc:PayableAmount></cac:LegalMonetaryTotal>"
            f"{''.join(kalemler)}</Invoice>")


NAKLIYECI = ("12345678901", "ÖRNEK NAKLİYE", "ÇANKAYA VD", "TCKN")
OTEL = ("2222222222", "DENEME TURİZM A.Ş.", "KUŞADASI VD")
ALICI_YURTDISI = ("2222222223", "SAMPLE GMBH", "")

DOSYALAR = {
    # 1) Tevkifatlı ticari fatura: 2 kalem, kalem iskontosu, 624 Yük taşımacılığı 2/10
    "tevkifat.xml": XML + fatura(
        "ABC2026000000001", "aaaaaaaa-0000-4000-8000-000000000001", "2026-09-05", "TICARIFATURA", "TEVKIFAT",
        NAKLIYECI, BIZ,
        [kalem(1, "Nakliye bedeli", 1, "C62", 26500, 20, tevkifat=("624", 2)),
         kalem(2, "Ek sefer", 2, "C62", 1000, 20, iskonto=100, tevkifat=("624", 2))],
        [vergi(28400, 5680, 20)], (28400, 100, 28400, 34080, 32944),
        tevkifat="<cac:WithholdingTaxTotal><cbc:TaxAmount currencyID=\"TRY\">1136</cbc:TaxAmount>"
                 + vergi(5680, 1136, 20, "624", "KDV TEVKIFAT") + "</cac:WithholdingTaxTotal>",
        ek="<cac:OrderReference><cbc:ID>SIP-77</cbc:ID><cbc:IssueDate>2026-09-01</cbc:IssueDate></cac:OrderReference>"
           "<cac:DespatchDocumentReference><cbc:ID>IRS2026000000010</cbc:ID><cbc:IssueDate>2026-09-04</cbc:IssueDate>"
           "</cac:DespatchDocumentReference>"),
    # 2) e-Arşiv, iki KDV oranı + konaklama vergisi
    "earsiv-otel.xml": XML + fatura(
        "GBB2026000000002", "aaaaaaaa-0000-4000-8000-000000000002", "2026-09-02", "EARSIVFATURA", "SATIS",
        OTEL, BIZ,
        [kalem(1, "Oda", 4, "DAY", 2000, 10, ek_vergi=vergi(8000, 160, 2, "0059", "KONAKLAMA VERGİSİ")),
         kalem(2, "Toplantı salonu", 1, "HUR", 1500, 20)],
        [vergi(8000, 800, 10), vergi(1500, 300, 20), vergi(8000, 160, 2, "0059", "KONAKLAMA VERGİSİ")],
        (9500, 0, 9500, 10760, 10760)),
    # 3) Dövizli ihracat istisna faturası (301)
    "ihracat-usd.xml": XML + fatura(
        "ABC2026000000003", "aaaaaaaa-0000-4000-8000-000000000003", "2026-09-23", "IHRACAT", "ISTISNA",
        BIZ, ALICI_YURTDISI,
        [kalem(1, "Alüminyum profil", 500, "KGM", 4.5, 0, para="USD")],
        [vergi(2250, 0, 0, istisna=("301", "11/1-a Mal İhracatı"), para="USD")], (2250, 0, 2250, 2250, 2250),
        para="USD", kur=41.25),
}
ZARF = (XML + '<sh:StandardBusinessDocument xmlns:sh="http://www.unece.org/cefact/namespaces/StandardBusinessDocumentHeader">'
        '<ef:Package xmlns:ef="http://www.efatura.gov.tr/package-namespace"><Elements><ElementList>'
        + DOSYALAR["tevkifat.xml"].replace(XML, "")
        + fatura("XYZ2026000000004", "aaaaaaaa-0000-4000-8000-000000000004", "2026-09-30", "TEMELFATURA", "IADE",
                 BIZ, OTEL, [kalem(1, "İade mal", 3, "MTR", 100, 20)], [vergi(300, 60, 20)], (300, 0, 300, 360, 360),
                 ek="<cac:BillingReference><cac:InvoiceDocumentReference><cbc:ID>GBB2026000000002</cbc:ID>"
                    "<cbc:IssueDate>2026-09-02</cbc:IssueDate></cac:InvoiceDocumentReference></cac:BillingReference>")
        + "</ElementList></Elements></ef:Package></sh:StandardBusinessDocument>")


def yaz(hedef: Path) -> None:
    hedef.mkdir(parents=True, exist_ok=True)
    for ad, metin in {**DOSYALAR, "zarf.xml": ZARF}.items():
        (hedef / ad).write_text(metin, encoding="utf-8")


def denetle(hedef: Path) -> None:
    from fatura_okuyucu.core import excel_yaz, isle
    belgeler = isle([hedef])
    by = {b.fatura_no: b for b in belgeler}
    assert len(belgeler) == 4, [b.fatura_no for b in belgeler]  # zarftaki kopya elenir
    t = by["ABC2026000000001"]
    assert (t.tutar, t.kdv, t.tevkifat, t.odenecek, t.vergiler_dahil) == (28400, 5680, 1136, 32944, 34080), t
    assert t.tevkifata_tabi_kdv == 5680 and t.tevkifata_tabi_tutar == 28400
    assert t.tevkifat_kodu == "624 2/10" and t.satici_vkn == "12345678901" and t.alici_vd == "MERKEZ VD"
    assert t.irsaliye_no == "IRS2026000000010" and t.siparis_no == "SIP-77" and t.kalem_sayisi == 2
    assert t.kalemler[1]["iskonto"] == 100 and t.kalemler[1]["iskonto_oran"] == 5 and t.kalemler[0]["birim"] == "ADET(UNIT)"
    o = by["GBB2026000000002"]
    assert (o.kdv, o.diger_vergi, o.odenecek) == (1100, 160, 10760) and o.senaryo == "e-Arşiv Fatura", o
    assert not o.uyarilar, o.uyarilar
    i = by["ABC2026000000003"]
    assert i.para_birimi == "USD" and i.kur == 41.25 and i.odenecek == round(2250 * 41.25, 2) and i.kdv == 0
    assert i.istisna == "301 11/1-a Mal İhracatı" and i.kalemler[0]["birim"] == "KİLOGRAM"
    r = by["XYZ2026000000004"]
    assert r.fatura_tipi == "IADE" and r.notlar.startswith("İade edilen fatura: GBB2026000000002")
    cikti = hedef / "deneme.xlsx"
    excel_yaz(belgeler, cikti)
    from openpyxl import load_workbook
    wb = load_workbook(cikti)
    assert wb.sheetnames[:3] == ["Faturalar", "Kalemler", "Vergiler"], wb.sheetnames
    assert wb["Kalemler"].max_row >= 7, wb["Kalemler"].max_row
    print("XML denetimi tamam:", len(belgeler), "fatura,", sum(len(b.kalemler) for b in belgeler), "kalem")


if __name__ == "__main__":
    yer = KOK / "tests" / "ornekler"
    yaz(yer)
    denetle(yer)
