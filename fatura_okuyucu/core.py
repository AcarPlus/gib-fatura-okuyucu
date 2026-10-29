"""Girdileri (RAR/ZIP/7z, klasör, PDF) toplar, PDF'leri okur, Excel'e yazar."""
import os
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .parsers import Belge, oku

ARSIV = {".rar", ".zip", ".7z"}


def _tar() -> str:
    sistem = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32" / "tar.exe"
    return str(sistem) if sistem.exists() else "tar"


def arsiv_ac(path: Path, hedef: Path) -> None:
    hedef.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                # Windows Gezgini ile yapılmış zip'lerde Türkçe adlar cp857 kodlu
                if not info.flag_bits & 0x800:
                    try:
                        info.filename = info.filename.encode("cp437").decode("cp857")
                    except UnicodeError:
                        pass
                z.extract(info, hedef)
        return
    # RAR / 7z: Windows'un kendi tar.exe'si (libarchive) açabiliyor
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    res = subprocess.run([_tar(), "-xf", str(path), "-C", str(hedef)],
                         capture_output=True, text=True, creationflags=flags)
    if res.returncode != 0:
        raise RuntimeError(f"{path.name} açılamadı: {res.stderr.strip() or res.returncode}")


def pdfleri_topla(girdiler: Iterable[Path], calisma: Path) -> list[Path]:
    """Girdilerdeki tüm PDF'leri (iç içe arşivler dahil) bulur."""
    pdfler: list[Path] = []

    def tara(kok: Path, derinlik: int) -> None:
        for p in sorted(kok.rglob("*")):
            if not p.is_file():
                continue
            if p.suffix.lower() == ".pdf":
                pdfler.append(p)
            elif p.suffix.lower() in ARSIV and derinlik < 3:
                ic = calisma / f"{len(list(calisma.iterdir()))}_{p.stem}"
                arsiv_ac(p, ic)
                tara(ic, derinlik + 1)

    for g in map(Path, girdiler):
        if g.is_dir():
            tara(g, 0)
        elif g.suffix.lower() == ".pdf":
            pdfler.append(g)
        elif g.suffix.lower() in ARSIV:
            ic = calisma / f"{len(list(calisma.iterdir()))}_{g.stem}"
            arsiv_ac(g, ic)
            tara(ic, 1)
    return pdfler


def isle(girdiler: Iterable[Path], ilerleme: Callable[[int, int, str], None] | None = None,
         iptal: Callable[[], bool] | None = None) -> list[Belge]:
    """Girdileri okuyup Belge listesi döndürür. ilerleme(i, toplam, dosya_adı)."""
    calisma = Path(tempfile.mkdtemp(prefix="fatura_okuyucu_"))
    try:
        pdfler = pdfleri_topla(girdiler, calisma)
        sonuc = []
        for i, p in enumerate(pdfler):
            if iptal and iptal():
                break
            if ilerleme:
                ilerleme(i, len(pdfler), p.name)
            b = oku(p)
            b.yol = "" if calisma in p.parents else str(p)  # geçici dosya yolu saklanmaz
            sonuc.append(b)
        if ilerleme:
            ilerleme(len(pdfler), len(pdfler), "")
        return sonuc
    finally:
        shutil.rmtree(calisma, ignore_errors=True)


# ---------------------------------------------------------------- Excel

# (başlık, Belge alanı, tür, genişlik)
SUTUNLAR = [
    ("DOSYA NO", "dosya", "metin", 38),
    ("FATURA NO", "fatura_no", "metin", 22),
    ("ALICI VERGİ DAİRESİ", "alici_vd", "metin", 28),
    ("ALICI VKN/TCKN", "alici_vkn", "metin", 15),
    ("SATICI VKN/TCKN", "satici_vkn", "metin", 15),
    ("TARİH", "tarih", "tarih", 12),
    ("TUTAR", "tutar", "para", 15),
    ("KDV", "kdv", "para", 13),
    ("ÖDENECEK TUTAR", "odenecek", "para", 16),
    # ek istek: tevkifat ve toplam alanları
    ("TOPLAM İSKONTO", "iskonto", "para", 13),
    ("KDV TEVKİFAT", "tevkifat", "para", 13),
    ("TEVKİFATA TABİ İŞLEM TUTARI", "tevkifata_tabi_tutar", "para", 16),
    ("TEVKİFATA TABİ İŞLEM ÜZ. HES. KDV", "tevkifata_tabi_kdv", "para", 16),
    ("VERGİLER DAHİL TOPLAM", "vergiler_dahil", "para", 16),
    ("DİĞER VERGİLER", "diger_vergi", "para", 13),
    ("G.V. STOPAJ", "stopaj", "para", 13),
    ("SATICI UNVANI", "satici_unvan", "metin", 32),
    ("ALICI UNVANI", "alici_unvan", "metin", 32),
    ("BELGE TÜRÜ", "tur", "metin", 22),
    ("KONTROL NOTU", "kontrol", "metin", 60),
]

_BASLIK = PatternFill("solid", fgColor="1F4E78")
_EK_BASLIK = PatternFill("solid", fgColor="2E75B6")
_BILGI_BASLIK = PatternFill("solid", fgColor="7F7F7F")
_SARI = PatternFill("solid", fgColor="FFF2CC")
_TOPLAM = PatternFill("solid", fgColor="DDEBF7")
_INCE = Side(style="thin", color="BFBFBF")


def excel_yaz(belgeler: list[Belge], hedef: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturalar"
    for c, (baslik, _, _, gen) in enumerate(SUTUNLAR, 1):
        h = ws.cell(1, c, baslik)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = _BASLIK if c <= 9 else _EK_BASLIK if c <= 16 else _BILGI_BASLIK
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = gen
    ws.row_dimensions[1].height = 32

    for r, b in enumerate(belgeler, 2):
        for c, (_, alan, tur, _) in enumerate(SUTUNLAR, 1):
            v = getattr(b, alan)
            cell = ws.cell(r, c, v if v not in ("",) else None)
            cell.border = Border(bottom=_INCE)
            if tur == "para":
                cell.number_format = "#,##0.00"
            elif tur == "tarih":
                cell.number_format = "DD.MM.YYYY"
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.number_format = "@"  # VKN'lerin baştaki sıfırı kaybolmasın
            if alan in b.eksik or (alan == "kontrol" and v):
                cell.fill = _SARI

    son = len(belgeler) + 1
    t = son + 1
    ws.cell(t, 1, "TOPLAM").font = Font(bold=True)
    for c, (_, _, tur, _) in enumerate(SUTUNLAR, 1):
        cell = ws.cell(t, c)
        cell.fill = _TOPLAM
        if tur == "para" and son >= 2:
            col = get_column_letter(c)
            cell.value = f"=SUBTOTAL(9,{col}2:{col}{son})"  # filtreyle birlikte çalışır
            cell.number_format = "#,##0.00"
            cell.font = Font(bold=True)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SUTUNLAR))}{max(son, 2)}"
    ws.sheet_view.zoomScale = 90

    a = wb.create_sheet("Açıklama")
    for satir in [
        ("GİB Fatura Okuyucu çıktısı", date.today().strftime("%d.%m.%Y")),
        ("Belge sayısı", len(belgeler)),
        ("Kontrol gereken belge", sum(1 for b in belgeler if b.eksik or b.kontrol)),
        (),
        ("Sarı hücre", "Bulunamadı ya da OCR ile okundu; PDF'le karşılaştırın"),
        ("TUTAR", "KDV hariç tutar (mal/hizmet toplamı, bilette toplam − KDV, makbuzda brüt ücret)"),
        ("ÖDENECEK TUTAR", "Tevkifat düşülmüş, ödenecek / tahsil edilen tutar"),
        ("Koyu mavi başlık", "İstenen sütunlar · açık mavi: ek tevkifat/toplam alanları · gri: bilgi"),
    ]:
        a.append(satir)
    a["A1"].font = Font(bold=True, size=13)
    a.column_dimensions["A"].width = 24
    a.column_dimensions["B"].width = 80

    hedef = Path(hedef)
    wb.save(hedef)
    return hedef
